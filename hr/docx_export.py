"""Word (.docx) version of an offer letter or payslip, so it can be edited in Word. Mirrors the on-screen paper (static/hr_docs.js)."""
import base64
import io
import os
import re
from datetime import datetime

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Mm, Pt, RGBColor

STATIC = os.path.join(os.path.dirname(__file__), "static")
MONTHS = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]
ONES = ["", "One", "Two", "Three", "Four", "Five", "Six", "Seven", "Eight", "Nine", "Ten", "Eleven", "Twelve", "Thirteen", "Fourteen", "Fifteen", "Sixteen", "Seventeen", "Eighteen", "Nineteen"]
TENS = ["", "", "Twenty", "Thirty", "Forty", "Fifty", "Sixty", "Seventy", "Eighty", "Ninety"]


def _two(n):
    return ONES[n] if n < 20 else TENS[n // 10] + ("-" + ONES[n % 10] if n % 10 else "")


def _three(n):
    return ONES[n // 100] + " Hundred" + (" " + _two(n % 100) if n % 100 else "") if n >= 100 else _two(n)


def words(n):
    n = int(abs(num(n)))
    if not n:
        return "Zero"
    parts = [(n // 10 ** 7, "Crore"), (n // 10 ** 5 % 100, "Lakh"), (n // 1000 % 100, "Thousand"), (n % 1000, "")]
    return " ".join((_two(v) + " " + w) if w else _three(v) for v, w in parts if v)


def num(v):
    try:
        return float(re.sub(r"[^0-9.\-]", "", str(v if v is not None else "")) or 0)
    except ValueError:
        return 0.0


def inr(n):
    return "₹{:,.2f}".format(num(n))


def _date(iso):
    try:
        return datetime.strptime(str(iso)[:10], "%Y-%m-%d")
    except ValueError:
        return None


def letter_date(iso):
    d = _date(iso)
    return f"{MONTHS[d.month - 1]} {d.day}, {d.year}" if d else ""


def slash_date(iso):
    d = _date(iso)
    return d.strftime("%d/%m/%Y") if d else ""


# ---- text
def all_vars(doc):
    c, s = doc.get("company") or {}, doc.get("signatory") or {}
    v = {"company_short": c.get("short_name") or c.get("name", ""), "company_legal": c.get("legal_name") or c.get("name", ""), "company_name": c.get("name", ""),
         "signatory_name": s.get("name", ""), "signatory_title": s.get("title", ""), "date": letter_date(doc.get("date"))}
    v.update({k: str(x) for k, x in (doc.get("vars") or {}).items() if x not in (None, "")})
    if v.get("salary") and not v.get("salary_words"):
        v["salary_words"] = words(num(v["salary"]))
    return v


def fill(text, v):
    return re.sub(r"\{\{\s*(\w+)\s*\}\}", lambda m: v.get(m.group(1), m.group(0)), str(text or ""))


def runs(par, text, v, size=None, bold=False, italic=False, underline=False):
    """Adds text to a paragraph, with **bold**, *italic* and line breaks."""
    text = fill(text, v)
    for tok in re.split(r"(\*\*.+?\*\*|\*.+?\*)", text, flags=re.S):
        if not tok:
            continue
        b, i = bold, italic
        if tok.startswith("**") and tok.endswith("**") and len(tok) > 4:
            tok, b = tok[2:-2], True
        elif tok.startswith("*") and tok.endswith("*") and len(tok) > 2:
            tok, i = tok[1:-1], True
        for n, line in enumerate(tok.split("\n")):
            r = par.add_run()
            if n:
                r.add_break()
            r.text = line
            r.bold, r.italic, r.underline = b, i, underline
            if size:
                r.font.size = Pt(size)


def _border(par, side="bottom", sz=8, color="111111", space=1):
    pPr = par._p.get_or_add_pPr()
    b = OxmlElement("w:pBdr")
    e = OxmlElement(f"w:{side}")
    for k, val in (("val", "single"), ("sz", str(sz)), ("space", str(space)), ("color", color)):
        e.set(qn(f"w:{k}"), val)
    b.append(e)
    pPr.append(b)


def _fmt(par, after=6, before=0, line=1.15, align=None):
    f = par.paragraph_format
    f.space_after, f.space_before, f.line_spacing = Pt(after), Pt(before), line
    if align is not None:
        par.alignment = align


def _logo(c):
    src = (c or {}).get("logo") or ""
    try:
        if src.startswith("data:image/"):
            return io.BytesIO(base64.b64decode(src.split(",", 1)[1]))
        if src.startswith("/hr/static/"):
            return open(os.path.join(STATIC, src[len("/hr/static/"):]), "rb")
    except (OSError, ValueError, IndexError):
        pass
    return None


def _cell_borders(cell, color="E1E3EA", sz=4):
    tcPr = cell._tc.get_or_add_tcPr()
    b = OxmlElement("w:tcBorders")
    for side in ("top", "left", "bottom", "right"):
        e = OxmlElement(f"w:{side}")
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), str(sz))
        e.set(qn("w:color"), color)
        b.append(e)
    tcPr.append(b)


def _shade(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear")
    sh.set(qn("w:color"), "auto")
    sh.set(qn("w:fill"), fill_hex.lstrip("#"))
    tcPr.append(sh)


def _no_borders(table):
    tblPr = table._tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{side}")
        e.set(qn("w:val"), "nil")
        b.append(e)
    tblPr.append(b)


# ---- offer letter
def _head_table(target, doc, v, size):
    """Letterhead as a two-column table in a header or in the body."""
    c = doc.get("company") or {}
    lh = doc.get("letterhead")
    if lh == "none":
        return
    table = target.add_table(1, 2, Mm(170)) if hasattr(target, "is_linked_to_previous") else target.add_table(1, 2)  # a page header needs a width
    _no_borders(table)
    left, right = table.rows[0].cells
    lp, rp = left.paragraphs[0], right.paragraphs[0]
    logo = _logo(c)
    if lh == "logo":
        runs(lp, doc.get("header_title") or "", v, size=size, bold=True, underline=True)
        if logo:
            rp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            rp.add_run().add_picture(logo, height=Inches(0.55))
    else:
        if logo:
            lp.add_run().add_picture(logo, height=Inches(0.55))
        rp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        lines = [("b", c.get("name", ""))] + [("t", "⚲ " + x) for x in str(c.get("address", "")).split("\n") if x.strip()]
        if c.get("phone"):
            lines.append(("t", "☎ " + c["phone"]))
        if c.get("email"):
            lines.append(("t", "✉ " + c["email"]))
        for n, (k, text) in enumerate(lines):
            r = rp.add_run()
            if n:
                r.add_break()
            r.text, r.bold = text, k == "b"
            r.font.size = Pt(size - 1.5)
    rule = target.add_paragraph()
    _border(rule, sz=24 if lh == "logo" else 8)
    _fmt(rule, after=10)
    return table


def offer_docx(data):
    d = Document()
    sec = d.sections[0]
    sec.page_width, sec.page_height = Mm(210), Mm(297)
    sec.left_margin = sec.right_margin = Mm(20)
    sec.top_margin, sec.bottom_margin = Mm(16), Mm(18)
    st = data.get("style") or {}
    font = (st.get("font") or "Times New Roman").split(",")[0].strip()
    size = float(st.get("size") or 11)
    normal = d.styles["Normal"]
    normal.font.name, normal.font.size = font, Pt(size)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), font)
    align = WD_ALIGN_PARAGRAPH.LEFT if st.get("align") == "left" else WD_ALIGN_PARAGRAPH.JUSTIFY
    v = all_vars(data)
    lh = data.get("letterhead")

    if lh == "logo":  # the same header on every page
        _head_table(sec.header, data, v, size)
        if sec.header.paragraphs and not sec.header.paragraphs[0].text and len(sec.header.paragraphs) > 1:
            sec.header.paragraphs[0]._p.getparent().remove(sec.header.paragraphs[0]._p)
    elif lh == "company":
        _head_table(d, data, v, size)

    p = d.add_paragraph("Date: " + letter_date(data.get("date")))
    _fmt(p, after=12, before=4)
    to = d.add_paragraph()
    for n, line in enumerate(str(data.get("recipient") or "").split("\n")):
        r = to.add_run()
        if n:
            r.add_break()
        r.text = line
    _fmt(to, after=14)

    count = 0
    for b in data.get("blocks") or []:
        t, text = b.get("t"), b.get("text", "")
        if t == "pb":
            d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        elif t == "h":
            p = d.add_paragraph()
            runs(p, text, v, bold=True, underline=True)
            _fmt(p, after=4, before=8)
        elif t == "num":
            count += 1
            p = d.add_paragraph()
            p.paragraph_format.left_indent, p.paragraph_format.first_line_indent = Inches(0.3), Inches(-0.3)
            p.add_run(f"{count}.\t")
            runs(p, text, v)
            _fmt(p, after=6, align=align)
        elif t == "sign":
            p = d.add_paragraph()
            runs(p, text, v)
            _fmt(p, after=6, before=12, align=WD_ALIGN_PARAGRAPH.LEFT)
        elif t == "accept":
            for label in [x for x in str(text).split("\n") if x.strip()]:
                p = d.add_paragraph()
                p.add_run(label + " " + "_" * 36)
                _fmt(p, after=14, before=6, align=WD_ALIGN_PARAGRAPH.LEFT)
        elif t == "sp":
            _fmt(d.add_paragraph(), after=10)
        else:
            p = d.add_paragraph()
            runs(p, text, v)
            _fmt(p, after=8, align=align)
    out = io.BytesIO()
    d.save(out)
    return out.getvalue()


# ---- payslip
def _tint(hex_color, amount=0.9):
    h = (hex_color or "#16a34a").lstrip("#")
    try:
        r, g, b = (int(h[i:i + 2], 16) for i in (0, 2, 4))
    except ValueError:
        r, g, b = 22, 163, 74
    return "%02X%02X%02X" % tuple(int(x + (255 - x) * amount) for x in (r, g, b))


def _txt(cell, text, bold=False, size=9, color=None, align=None, after=2):
    p = cell.paragraphs[0] if not cell.paragraphs[0].text and not cell.paragraphs[0].runs else cell.add_paragraph()
    r = p.add_run(str(text))
    r.bold, r.font.size = bold, Pt(size)
    if color:
        r.font.color.rgb = RGBColor.from_string(color)
    if align is not None:
        p.alignment = align
    _fmt(p, after=after, line=1.05)
    return p


def payslip_docx(data):
    d = Document()
    sec = d.sections[0]
    sec.page_width, sec.page_height = Mm(210), Mm(297)
    sec.left_margin = sec.right_margin = Mm(18)
    sec.top_margin, sec.bottom_margin = Mm(16), Mm(16)
    normal = d.styles["Normal"]
    normal.font.name, normal.font.size = "Arial", Pt(9)
    c = data.get("company") or {}
    tint = _tint(data.get("accent"))
    R = WD_ALIGN_PARAGRAPH.RIGHT

    head = d.add_table(1, 2)
    _no_borders(head)
    a, b = head.rows[0].cells
    _txt(a, c.get("name", ""), bold=True, size=13)
    _txt(a, str(c.get("address", "")).replace("\n", ", "), size=8, color="6B7280")
    _txt(b, "Payslip For the Month", size=8, color="6B7280", align=R)
    _txt(b, data.get("period", ""), bold=True, size=11, align=R)
    rule = d.add_paragraph()
    _border(rule, sz=4, color="D9D9DE")

    t = d.add_table(1, 2)
    _no_borders(t)
    left, right = t.rows[0].cells
    _txt(left, "EMPLOYEE SUMMARY", bold=True, size=8, color="4B5563", after=6)
    for r in data.get("rows") or []:
        p = _txt(left, "", size=9)
        x = p.add_run(f"{r.get('label', '')}\t:  ")
        x.font.size, x.font.color.rgb = Pt(9), RGBColor.from_string("6B7280")
        y = p.add_run(str(r.get("value", "")))
        y.font.size = Pt(9)
        p.paragraph_format.tab_stops.add_tab_stop(Inches(1.35))
    gross = sum(num(x.get("amount")) for x in data.get("earnings") or [])
    ded = sum(num(x.get("amount")) for x in data.get("deductions") or [])
    net = gross - ded
    _shade(right, tint)
    _cell_borders(right)
    _txt(right, inr(net), size=18)
    _txt(right, "Total Net Pay", size=8.5, color="4B5563", after=8)
    _txt(right, f"Paid Days\t:  {data.get('paid_days', '')}", size=9)
    _txt(right, f"LOP Days\t:  {data.get('lop_days', '')}", size=9)
    _fmt(d.add_paragraph(), after=4)

    e, dd = data.get("earnings") or [], data.get("deductions") or []
    n = max(len(e), len(dd))
    table = d.add_table(n + 2, 4)
    for row in table.rows:
        for cell in row.cells:
            _cell_borders(cell)
    for i, h in enumerate(("EARNINGS", "AMOUNT", "DEDUCTIONS", "AMOUNT")):
        _txt(table.cell(0, i), h, bold=True, size=8.5, align=R if i % 2 else None)
    for i in range(n):
        for col, items in ((0, e), (2, dd)):
            if i < len(items):
                _txt(table.cell(i + 1, col), items[i].get("label", ""), size=9)
                _txt(table.cell(i + 1, col + 1), inr(items[i].get("amount")), bold=True, size=9, align=R)
    for col, label, val in ((0, "Gross Earnings", gross), (2, "Total Deductions", ded)):
        _txt(table.cell(n + 1, col), label, bold=True, size=9)
        _txt(table.cell(n + 1, col + 1), inr(val), bold=True, size=9, align=R)
        _shade(table.cell(n + 1, col), "F6F6FA")
        _shade(table.cell(n + 1, col + 1), "F6F6FA")
    _fmt(d.add_paragraph(), after=4)

    pay = d.add_table(1, 2)
    for cell in pay.rows[0].cells:
        _cell_borders(cell)
    _txt(pay.cell(0, 0), "TOTAL NET PAYABLE", bold=True, size=9)
    _txt(pay.cell(0, 0), "Gross Earnings - Total Deductions", size=8, color="6B7280")
    _txt(pay.cell(0, 1), inr(net), bold=True, size=11, align=R)
    _shade(pay.cell(0, 1), tint)
    if data.get("show_words"):
        p = d.add_paragraph()
        p.alignment = R
        x = p.add_run("Amount In Words : ")
        x.font.size, x.font.color.rgb = Pt(8.5), RGBColor.from_string("6B7280")
        y = p.add_run(f"Indian Rupee {words(net)} Only")
        y.font.size = Pt(8.5)
        _fmt(p, before=8)
    foot = d.add_paragraph(data.get("footer", ""))
    foot.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for r in foot.runs:
        r.font.size, r.font.color.rgb = Pt(7.5), RGBColor.from_string("6B7280")
    _fmt(foot, before=16)
    out = io.BytesIO()
    d.save(out)
    return out.getvalue()


def build(data):
    return payslip_docx(data) if data.get("kind") == "payslip" else offer_docx(data)
